"""Import parsers smoke test: CSV, XLSX, DOCX, PDF and JSON."""
import csv
import tempfile
import sys
from pathlib import Path

from openpyxl import Workbook
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from importador import parse_document

root = Path(tempfile.mkdtemp(prefix="libris-import-test-"))

# CSV with common Portuguese headings and semicolon separator.
csv_path = root / "acervo.csv"
with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.writer(handle, delimiter=";")
    writer.writerow(["Título", "Autor", "ISBN", "Gênero", "Coluna", "Fileira", "Tombo"])
    writer.writerow(["O Pequeno Príncipe", "Antoine de Saint-Exupéry", "9780000000001", "Fábula", "A", "2", "ESC-001"])
books, warnings = parse_document(csv_path, "books")
assert books and books[0]["title"] == "O Pequeno Príncipe", (books, warnings)
assert books[0]["accession_code"] == "ESC-001" and books[0]["shelf_row"] == "2"

# XLSX student roster, with class and no school registration (auto-registration allowed at import).
xlsx_path = root / "alunos.xlsx"
wb = Workbook()
ws = wb.active
ws.append(["Nome completo", "Turma", "Telefone", "Meta anual"])
ws.append(["Ana Beatriz", "5º ano A", "81999990000", 10])
wb.save(xlsx_path)
readers, warnings = parse_document(xlsx_path, "readers")
assert readers and readers[0]["name"] == "Ana Beatriz", (readers, warnings)
assert readers[0]["class_name"] == "5º ano A" and readers[0]["reading_goal"] == "10"

# DOCX table for a book catalog.
docx_path = root / "leitura.docx"
document = Document()
table = document.add_table(rows=2, cols=3)
table.style = "Table Grid"
for i, value in enumerate(["Nome do Livro", "Escritor", "Categoria"]):
    table.cell(0, i).text = value
for i, value in enumerate(["Quarto de Despejo", "Carolina Maria de Jesus", "Memórias"]):
    table.cell(1, i).text = value
document.save(docx_path)
doc_books, warnings = parse_document(docx_path, "books")
assert doc_books and doc_books[0]["title"] == "Quarto de Despejo", (doc_books, warnings)
assert doc_books[0]["author"] == "Carolina Maria de Jesus"

# Searchable PDF with a ruled table.
pdf_path = root / "catalogo.pdf"
pdf_doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
data = [
    ["Título", "Autor", "ISBN", "Gênero"],
    ["Capitães da Areia", "Jorge Amado", "9780000000002", "Romance"],
]
pdf_table = Table(data, colWidths=[180, 140, 120, 100])
pdf_table.setStyle(TableStyle([
    ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
]))
pdf_doc.build([pdf_table])
pdf_books, warnings = parse_document(pdf_path, "books")
assert pdf_books and pdf_books[0]["title"] == "Capitães da Areia", (pdf_books, warnings)
assert pdf_books[0]["author"] == "Jorge Amado"

# JSON export.
json_path = root / "students.json"
json_path.write_text('[{"Nome":"Pedro Santos","Matrícula":"ALU-992","Turma":"7º ano"}]', encoding="utf-8")
json_readers, warnings = parse_document(json_path, "readers")
assert json_readers and json_readers[0]["name"] == "Pedro Santos"
assert json_readers[0]["registration"] == "ALU-992"

print("OK: importação de CSV, XLSX, DOCX, PDF pesquisável e JSON validada.")
